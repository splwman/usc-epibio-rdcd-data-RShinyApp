from __future__ import annotations

import json
import os
import re
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK_PATH = Path("RDCD_Data_Source_Preparation_Template.xlsx")
SHEET_NAME = "Data_Prep"

# Maps the Form/Power Automate field names to the actual public catalog columns.
# One Form value may populate more than one catalog column when that improves
# the app's display and filtering.
FIELD_MAP = {
    "record_id": ("record_id",),
    "data_source": ("data_source", "official_link_text"),
    "official_link": ("official_url",),
    "narrative": ("main_data_type_use",),
    "access_level": ("raw_access_level", "access_level_standardized"),
    "primary_category": ("primary_category",),
    "secondary_categories": ("secondary_categories_tags",),
    "data_type": ("data_type_standardized",),
    "geography": ("geography_coverage",),
    "population": ("population",),
    "unit_of_analysis": ("unit_of_analysis",),
    "keywords": ("keywords_for_search",),
    "linkage_potential": ("linkage_potential",),
    "access_burden": ("access_burden",),
    "rdcd_notes": ("rdcd_consultation_notes",),
}
REQUIRED_WORKBOOK_COLUMNS = ("record_id", "data_source")


def normalize_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def safe_text(value: object) -> str:
    text = "" if value is None else str(value).strip()
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def find_header_row(worksheet) -> tuple[int, dict[str, int]]:
    for row_number in range(1, min(worksheet.max_row, 20) + 1):
        headers = {
            normalize_header(cell.value): cell.column
            for cell in worksheet[row_number]
            if normalize_header(cell.value)
        }
        if "record_id" in headers and "data_source" in headers:
            return row_number, headers

    raise ValueError(
        "Could not find a header row containing both 'record_id' and 'data_source' "
        "within the first 20 rows of the Data_Prep sheet."
    )


def copy_row_style(worksheet, source_row: int, target_row: int, columns: list[int]) -> None:
    for column in columns:
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)


def find_matching_table(worksheet, header_row: int):
    for table in worksheet.tables.values():
        start_cell, _ = table.ref.split(":")
        if worksheet[start_cell].row == header_row:
            return table

    raise ValueError("Could not find an Excel Table starting on the header row.")


def expand_table(table, target_row: int) -> None:
    start_cell, end_cell = table.ref.split(":")
    end_column = re.sub(r"\d+$", "", end_cell)
    table.ref = f"{start_cell}:{end_column}{target_row}"


def main() -> None:
    payload = json.loads(os.environ["SUBMISSION_JSON"])

    record_id = safe_text(payload.get("record_id"))
    if not re.fullmatch(r"DS-\d{3,}", record_id):
        raise ValueError("record_id must use the form DS-001, DS-002, and so on.")

    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    workbook = load_workbook(WORKBOOK_PATH)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Worksheet not found: {SHEET_NAME}")

    worksheet = workbook[SHEET_NAME]
    header_row, headers = find_header_row(worksheet)

    missing_columns = [
        column for column in REQUIRED_WORKBOOK_COLUMNS if column not in headers
    ]
    if missing_columns:
        raise ValueError(
            "Workbook is missing required columns: " + ", ".join(missing_columns)
        )

    existing_ids = {
        safe_text(worksheet.cell(row, headers["record_id"]).value)
        for row in range(header_row + 1, worksheet.max_row + 1)
    }
    if record_id in existing_ids:
        raise ValueError(f"record_id already exists in the workbook: {record_id}")

    table = find_matching_table(worksheet, header_row)
    _, table_end_cell = table.ref.split(":")
    target_row = worksheet[table_end_cell].row + 1
    mapped_columns = sorted(
        {
            headers[catalog_column]
            for catalog_columns in FIELD_MAP.values()
            for catalog_column in catalog_columns
            if catalog_column in headers
        }
    )
    copy_row_style(worksheet, target_row - 1, target_row, mapped_columns)

    for form_field, catalog_columns in FIELD_MAP.items():
        value = safe_text(payload.get(form_field))
        for catalog_column in catalog_columns:
            if catalog_column in headers:
                worksheet.cell(target_row, headers[catalog_column]).value = value

    expand_table(table, target_row)
    workbook.save(WORKBOOK_PATH)

    summary_lines = [
        f"# Review data-source submission: {record_id}",
        "",
        "Please verify the proposed source information before merging.",
        "",
    ]
    for form_field in FIELD_MAP:
        if form_field == "record_id":
            continue
        label = form_field.replace("_", " ").title()
        summary_lines.append(f"**{label}:** {safe_text(payload.get(form_field)) or '—'}")
        summary_lines.append("")

    Path("/tmp/form_submission_review.md").write_text(
        "\n".join(summary_lines), encoding="utf-8"
    )


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"Submission update failed: {error}", file=sys.stderr)
        raise
